import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.styles import Border, Side
from openpyxl.styles import PatternFill
df = pd.read_excel("students.xlsx")
df.columns = df.columns.str.strip()
df = df.dropna(how="all")
required_columns = ["Name", "Physics", "Chemistry", "Biology"]
for column in required_columns:
    if column not in df.columns:
        print(f"Missing column: {column}")
        exit()
df["Total"] = df["Physics"] + df["Chemistry"] + df["Biology"]
df["Percentage"] = df["Total"] / 3
def calculate_grade(percentage):
    if percentage >= 90:
        return "A"
    elif percentage >= 75:
        return "B"
    elif percentage >= 50:
        return "C"
    else:
        return "Fail"
def check_status(percentage):
	if percentage>=50:
		return "Pass"
	else:
		return "Fail"
df["Grade"] = df["Percentage"].apply(calculate_grade)
df["Status"] = df["Percentage"].apply(check_status)
df["Rank"] = df["Percentage"].rank(ascending=False, method="min").astype(int)
topper_name = df.loc[df["Percentage"].idxmax(), "Name"]
df["Topper"] = df["Name"].apply(
    lambda name: "Yes" if name == topper_name else "No")
df = df.sort_values("Rank")
df["Percentage"] = df["Percentage"].round(2)
df.to_excel("student_report.xlsx", index=False)
workbook = load_workbook("student_report.xlsx")
sheet = workbook.active
# Format header
for cell in sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

# Adjust column widths
for column in sheet.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    sheet.column_dimensions[column_letter].width = max_length + 2
# Format percentage column
for cell in sheet["F"][1:]:
    cell.number_format = "0.00"
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"))
for row in sheet.iter_rows():
    for cell in row:
        cell.border = thin_border
 # Freeze header row
sheet.freeze_panes = "A2"

# Add filter
sheet.auto_filter.ref = sheet.dimensions
# Highlight topper and failed students

topper_fill = PatternFill(fill_type="solid", fgColor="00FF00")
fail_fill = PatternFill(fill_type="solid", fgColor="FF0000")

for row in range(2, sheet.max_row + 1):

    if sheet[f"J{row}"].value == "Yes":
        for cell in sheet[row]:
            cell.fill = topper_fill
            cell.font = Font(bold=True)

    elif sheet[f"H{row}"].value == "Fail":
        for cell in sheet[row]:
            cell.fill = fail_fill
            cell.font = Font(bold=True)
# Create summary
total_students = len(df)
passed_students = (df["Status"] == "Pass").sum()
failed_students = (df["Status"] == "Fail").sum()
average_percentage = df["Percentage"].mean()
highest_percentage = df["Percentage"].max()

summary_start = total_students + 3  # +1 for header row, +2 for gap

sheet[f"A{summary_start}"] = "REPORT SUMMARY"
sheet[f"A{summary_start+1}"] = "Total Students"
sheet[f"B{summary_start+1}"] = total_students

sheet[f"A{summary_start+2}"] = "Passed"
sheet[f"B{summary_start+2}"] = passed_students

sheet[f"A{summary_start+3}"] = "Failed"
sheet[f"B{summary_start+3}"] = failed_students

sheet[f"A{summary_start+4}"] = "Average Percentage"
sheet[f"B{summary_start+4}"] = round(average_percentage, 2)

sheet[f"A{summary_start+5}"] = "Highest Percentage"
sheet[f"B{summary_start+5}"] = highest_percentage

sheet[f"A{summary_start+6}"] = "Topper"
sheet[f"B{summary_start+6}"] = topper_name

# Format summary
for row in range(summary_start, summary_start+7):
    for col in range(1, 3):
        cell = sheet.cell(row=row, column=col)
        cell.border = thin_border

sheet[f"A{summary_start}"].font = Font(bold=True)
sheet[f"A{summary_start}"].alignment = Alignment(horizontal="center")

for row in range(summary_start+1, summary_start+7):
    sheet[f"A{row}"].font = Font(bold=True)

sheet.column_dimensions["A"].width = 20
sheet.column_dimensions["B"].width = 18

# Style summary heading
summary_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")

for cell in sheet[summary_start][0:2]:
    cell.fill = summary_fill
    cell.font = Font(bold=True)

workbook.save("automated_student_report.xlsx")
